## 惠州LNG技能题库word题库格式检验及excel表生成
import docx
import openpyxl
import re

def extract_question_bank(input_file, output_file):
    # 创建Excel工作簿
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    # 设置表头
    headers = ["试题难度（容易、较易、中等、较难、很难）","试题类型", "试题内容", "正确答案",
               "题目选项A", "题目选项B", "题目选项C", "题目选项D", "题目选项E", "题目选项F"]
    sheet.append(headers)

    # 打开Word文件
    doc = docx.Document(input_file)

    # 提取题目内容并填充Excel表格
    content = ""
    for para in doc.paragraphs:
        content += para.text

    # 使用正则表达式匹配题目内容
    pattern = r"(\d+、L\d+)(.*?)((?=\d+、L\d+)|$)"
    matches = re.findall(pattern, content, re.DOTALL)

    # 题目信息
    i = 0
    j = -1
    for match in matches:
        i = i+1
        difficulty = extract_difficulty(match[0])
        question_content = match[1]
        pre_question_content = matches[i-2][1]

        question_type, correct_answer, options = parse_question(question_content)
        pre_question_type, pre_correct_answer, pre_options = parse_question(pre_question_content)

        # 系统匹配
        # patternsys = r"(##[\u4e00-\u9fa5]+##)"
        patternsys = r"(##.{3,20}##)"
        matchessys = re.findall(patternsys, content, re.DOTALL)
        if pre_question_type == "简答题" and question_type == "单选题":
            j = j+1
            if j == 13:
                j=12
        sys = matchessys[j]
        sys = sys.replace("##","")

        
        if question_type == "单选题":
            question_content = question_content.replace(options[0],'')
        elif question_type == "多选题":
            question_content = question_content.replace(options[0],'')
        # elif question_type == "简答题":
        #     question_content = question_content.replace('答：','')
        question_content = question_content.replace(correct_answer,'')
        # 将题目信息填充到Excel表格中
        add_question_to_sheet(sheet, difficulty, sys, question_content, question_type, correct_answer, *options)

    # 保存Excel文件
    workbook.save(output_file)

def extract_difficulty(difficulty_str):
    # 去除 "、" 及其前面的字符
    return difficulty_str.split("、")[-1]

for i in range(12):
    def parse_question(question_content):
        # 判断题型和答案
        if "（对）" in question_content or "(对)" in question_content or "【对】" in question_content:
            question_type = "判断题"
            correct_answer = "对"
            options = []
        elif "（错）" in question_content or "(错)" in question_content or "【错】" in question_content:
            question_type = "判断题"
            correct_answer = "错"
            options = []
        elif "答：" in question_content:
            question_type = "简答题"
            split_index = question_content.index("答：")
            correct_answer = question_content[split_index:].strip()
            options = []
        else:
            options = extract_options(question_content)
        
            option_count = len(options)
            if option_count == 1:
                question_type = "单选题"
                correct_answer = extract_correct_answer(question_content)
                options[0] = options[0][options[0].index("A"):]
                # 判断中文括号（）内的字母数量
                chinese_brackets = re.findall(r"（([A-F]+)）", question_content)
                if len(correct_answer) > 1:
                    question_type = "多选题"
            else:
                question_type = ""
                correct_answer = ""
                options = []

        return question_type, correct_answer, options

def extract_options(question_content):
    # 通过正则表达式提取选项
    options = re.findall(r"（[A-F]+）(.*?)(?=(?:（[A-F]）+)|$)", question_content)
    if options == []:
        ff = "opps"
    return options

def extract_correct_answer(question_content):
    # 通过正则表达式提取括号中的字母
    match = re.search(r"（([A-F]+)）", question_content)
    if match:
        return match.group(1)
    else:
        return ""

def add_question_to_sheet(sheet, difficulty, sys, question_content, question_type, correct_answer, *options):
    # 填充题目信息到Excel表格中
    row_data = [difficulty, sys, question_type, question_content, correct_answer]

    if question_type == "单选题":
        optionA = re.findall(r'A、(.*)B、',options[0])
        optionB = re.findall(r'B、(.*)C、',options[0])
        optionC = re.findall(r'C、(.*)D、',options[0])
        optionD = re.findall(r'D、(.*)$',options[0])
        
        row_data.extend(optionA)
        row_data.extend(optionB)
        row_data.extend(optionC)
        row_data.extend(optionD)
        
    elif question_type == "多选题":
        optionA = re.findall(r'A、(.*)B、',options[0])
        optionB = re.findall(r'B、(.*)C、',options[0])
        optionC = re.findall(r'C、(.*)D、',options[0])
        optionD = re.findall(r'D、(.*)$',options[0])

        if  re.findall(r'E、(.*)F',options[0]):
            
            optionD = re.findall(r'D、(.*)E',options[0])
            optionE = re.findall(r'E、(.*)F',options[0])
            optionF = re.findall(r'F、(.*)$',options[0])
            ff = 'test'
        elif  re.findall(r'E、(.*)$',options[0]):
            optionD = re.findall(r'D、(.*)E',options[0])
            optionE = re.findall(r'E、(.*)$',options[0])
            optionF = ''
            ff = 'test'
        else:
            optionE = ''
            optionF = ''
        row_data.extend(optionA)
        row_data.extend(optionB)
        row_data.extend(optionC)
        row_data.extend(optionD)
        row_data.extend(optionE)
        row_data.extend(optionF)
    else: 
        row_data.extend(options)
    sheet.append(row_data)

# 输入和输出文件路径
input_file = "input.docx"
output_file = "output.xlsx"

# 提取题库并生成Excel文件
extract_question_bank(input_file, output_file)